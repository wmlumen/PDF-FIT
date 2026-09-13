import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

doc = Document()

# Estilos globales
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(11)

# Título Memorándum
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Memorándum")
run.bold = True
run.font.size = Pt(14)
run.underline = True

# Encabezado
p_header = doc.add_paragraph()
p_header.add_run("A: ").bold = True
p_header.add_run("Mg. Luz María Benegas de Guevara - Dirección General – Mg. Lourdes Jaqueline Cáceres- Dirección Académica y alumnos/as\n")
p_header.add_run("De: ").bold = True
p_header.add_run("Lic. Cristian Kein\n")
p_header.add_run("Fecha: ").bold = True
p_header.add_run("04/09/2026\n")
p_header.add_run("Objeto: ").bold = True
p_header.add_run("Remitir cronograma de clases del mes de septiembre")

doc.add_paragraph("_" * 70)

# Datos de la Asignatura
p_datos = doc.add_paragraph()
p_datos.add_run("CARRERA: ").bold = True
p_datos.add_run("Administración Aduanera – Administración de Empresas – Gestión Pública\n")
p_datos.add_run("Sección: ").bold = True
p_datos.add_run("S  - ")
p_datos.add_run("Hora: ").bold = True
p_datos.add_run("13:00 a 15:40\n")
p_datos.add_run("ASIGNATURA: ").bold = True
p_datos.add_run("Tecnología de la Información y Comunicación (TIC)\n")
p_datos.add_run("MES Y AÑO: ").bold = True
p_datos.add_run("Septiembre 2026")

doc.add_paragraph()

# Tabla de Cronograma
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'CLASES'
hdr_cells[1].text = 'FECHA'
hdr_cells[2].text = 'HORA'

for cell in hdr_cells:
    paragraphs = cell.paragraphs
    for paragraph in paragraphs:
        for run in paragraph.runs:
            run.font.bold = True
            run.font.name = 'Arial'

datos_tabla = [
    ("Clase presencial física Sábado", "05-09-2026", "13:00 a 15:40"),
    ("Clase virtual presencial sincrónica (Miércoles)", "09-09-2026", "13:00 a 15:40"),
    ("Clase presencial física Sábado", "12-09-2026", "13:00 a 15:40"),
    ("Clase virtual presencial sincrónica (Miércoles)", "16-09-2026", "13:00 a 15:40"),
    ("Clase presencial física Sábado", "19-09-2026", "13:00 a 15:40"),
    ("Clase virtual presencial sincrónica (Miércoles)", "30-09-2026", "13:00 a 15:40"),
    ("Entrega de trabajo y examen final (Sábado)", "03-10-2026", "13:00 a 15:40")
]

for clase, fecha, hora in datos_tabla:
    row_cells = table.add_row().cells
    row_cells[0].text = clase
    row_cells[1].text = fecha
    row_cells[2].text = hora

doc.add_paragraph()

# Nota asincrónica para claridad de Dirección Académica
p_obs = doc.add_paragraph()
p_obs.add_run("Observación Académica: ").bold = True
p_obs.add_run("Del Miércoles 23 al Martes 29 de Septiembre se desarrollará la 'Semana Asincrónica' donde los alumnos completarán trabajos prácticos en la plataforma sin encuentro sincrónico en vivo, justificando la ausencia de encuentros en dichas fechas.")
p_obs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p_obs.style.font.size = Pt(10)

doc.add_paragraph()

# Recordatorio final
p_recordatorio = doc.add_paragraph()
run_rec = p_recordatorio.add_run("RECORDAR DE PRESENTAR EL INSTRUMENTO DE EVALUACIÓN EL 28 DE SEPTIEMBRE DEL 2026 PARA ANÁLISIS Y CONFIRMACIÓN POR MESA DE ENTRADA.")
run_rec.bold = True
run_rec.font.color.rgb = RGBColor(255, 0, 0) # Rojo para destacar
p_recordatorio.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Guardar
output_path = os.path.join(os.getcwd(), "Memorandum_Cronograma_TIC.docx")
doc.save(output_path)
print(f"Memorándum guardado en {output_path}")
