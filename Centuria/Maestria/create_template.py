import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = docx.Document()
# Set margins to 1 inch (APA 7)
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Cover Page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("INSTITUTO SUPERIOR EN CIENCIAS EMPRESARIALES\n“CENTURIA”\n\nTRABAJO DE INVESTIGACIÓN\nMONOGRAFÍA\n")
run.bold = True
run.font.name = 'Times New Roman'
run.font.size = Pt(14)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
run = p.add_run("\nMATERIA: {{ materia }}\nDOCENTE: {{ docente }}\nALUMNO: {{ alumno }}\nCEDULA DE IDENTIDAD N°: {{ cedula }}\nCARRERA: {{ carrera }}\nSECCION / GRUPO: {{ seccion }}\n")
run.font.name = 'Times New Roman'
run.font.size = Pt(12)

# Institutional Footer / Table
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("\nDATOS PARA USO INTERNO\n")
run.bold = True; run.font.name = 'Times New Roman'; run.font.size = Pt(12)

table = doc.add_table(rows=2, cols=3)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'ADMINISTRACION'
hdr_cells[1].text = 'COORDINACION'
hdr_cells[2].text = 'DOCENTE'
row_cells = table.rows[1].cells
row_cells[0].text = '\n\n\nFIRMA Y SELLO'
row_cells[1].text = '\n\n\nFIRMA Y SELLO'
row_cells[2].text = '\n\n\nFIRMA Y SELLO'
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)

doc.add_page_break()

# Indices
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("ÍNDICE GENERAL")
run.bold = True; run.font.name = 'Times New Roman'; run.font.size = Pt(12)

p_toc = doc.add_paragraph()
run_toc = p_toc.add_run()
fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'separate')
fldChar3 = OxmlElement('w:fldChar'); fldChar3.set(qn('w:fldCharType'), 'end')
run_toc._r.append(fldChar1); run_toc._r.append(instrText); run_toc._r.append(fldChar2); run_toc._r.append(fldChar3)

doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("ÍNDICE DE TABLAS")
run.bold = True; run.font.name = 'Times New Roman'; run.font.size = Pt(12)

p_toc2 = doc.add_paragraph()
run_toc2 = p_toc2.add_run()
fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'TOC \\h \\z \\c "Tabla"'
fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'separate')
fldChar3 = OxmlElement('w:fldChar'); fldChar3.set(qn('w:fldCharType'), 'end')
run_toc2._r.append(fldChar1); run_toc2._r.append(instrText); run_toc2._r.append(fldChar2); run_toc2._r.append(fldChar3)

doc.add_page_break()
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run("ÍNDICE DE FIGURAS")
run.bold = True; run.font.name = 'Times New Roman'; run.font.size = Pt(12)

p_toc3 = doc.add_paragraph()
run_toc3 = p_toc3.add_run()
fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'TOC \\h \\z \\c "Figura"'
fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'separate')
fldChar3 = OxmlElement('w:fldChar'); fldChar3.set(qn('w:fldCharType'), 'end')
run_toc3._r.append(fldChar1); run_toc3._r.append(instrText); run_toc3._r.append(fldChar2); run_toc3._r.append(fldChar3)

doc.add_page_break()
# Jinja2 tag para docxtpl
p = doc.add_paragraph('{{ body_content }}')

doc.save('Plantilla_Centuria_APA7.docx')
print("Template created.")
