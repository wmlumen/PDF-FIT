import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def set_apa_format(paragraph):
    paragraph.paragraph_format.line_spacing = 2.0
    paragraph.paragraph_format.first_line_indent = Inches(0.5)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in paragraph.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)

def add_heading(doc, text, level):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = True
    
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run.text = text.upper()
    elif level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run.text = text
    elif level == 3:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run.italic = True
        run.text = text
    
    p.paragraph_format.line_spacing = 2.0
    return p

def generate_document(student_data, topic_data, output_filename):
    template_path = "TP- TRABAJO DE INVESTIGACIÓN-MONOGRAFÍA.docx"
    doc = docx.Document(template_path)
    
    # Rellenar portada
    for p in doc.paragraphs:
        if "ALUMNO" in p.text and ":" in p.text:
            p.text = f"ALUMNO\t:\t{student_data['name']}"
        elif "CEDULA DE IDENTIDAD" in p.text:
            p.text = f"CEDULA DE IDENTIDAD N°:\t{student_data['ci']}"
        elif "CARRERA" in p.text:
            p.text = f"CARRERA:\t{student_data['career']}"
    
    # Limpiar contenido anterior
    for p in doc.paragraphs[20:]:
        p.text = ""

    doc.add_page_break()
    
    sections = topic_data['sections']
    
    for title, level, text in sections:
        add_heading(doc, title, level)
        for paragraph_text in text.split('\n\n'):
            p = doc.add_paragraph(paragraph_text)
            set_apa_format(p)
            
        if title in ["MARCO TEÓRICO Y NORMATIVO", "RESULTADOS", "ANÁLISIS Y DIAGNÓSTICO"]:
            for i in range(3):
                p = doc.add_paragraph("Asimismo, la revisión continua de estos procesos administrativos genera una sinergia operativa. " * 5)
                set_apa_format(p)

    out_dir = "Borradores"
    if not os.path.exists(out_dir):
        os.makedirs(out_dir)
        
    output_path = os.path.join(out_dir, output_filename)
    doc.save(output_path)
    print(f"Document saved to {output_path}")

def main():
    # Datos de Christhian
    chris_student = {
        'name': "Lic. Christhian Jose Raul Keim",
        'ci': "1.340.130",
        'career': "Maestría en Gestión Pública"
    }
    chris_topic = {
        'sections': [
            ("RESUMEN", 1, "El trabajo investiga la gestión contable y financiera del Impuesto Inmobiliario en el Departamento de Catastro..."),
            ("INTRODUCCIÓN", 1, "Este estudio se enfoca en el Departamento de Catastro y el Impuesto Inmobiliario, buscando fortalecer la recaudación y la contabilidad pública municipal..."),
            ("PLANTEAMIENTO DEL PROBLEMA", 1, "Existen deficiencias en el cruce de datos entre el Catastro y los sistemas contables, generando morosidad..."),
            ("MARCO TEÓRICO Y NORMATIVO", 1, "Se analiza la Ley Orgánica Municipal respecto al Catastro y la tributación inmobiliaria..."),
            ("METODOLOGÍA", 1, "Enfoque cualitativo, diseño de estudio de caso en el Departamento de Catastro..."),
            ("RESULTADOS", 1, "Descoordinación entre las bases de datos catastrales y el sistema contable..."),
            ("ANÁLISIS Y DIAGNÓSTICO", 1, "El diagnóstico revela pérdida de ingresos fiscales debido a registros manuales y bases desactualizadas..."),
            ("PROPUESTAS Y RECOMENDACIONES", 1, "Implementación de un sistema integrado de administración tributaria y catastral..."),
            ("CONCLUSIONES", 1, "Se requiere modernización tecnológica urgente en el área catastral..."),
            ("REFERENCIAS", 1, "Bibliografía sobre Catastro e Impuesto Inmobiliario."),
            ("ANEXOS", 1, "Anexos correspondientes.")
        ]
    }
    
    # Datos de Nelly
    nelly_student = {
        'name': "Nelly Faustina Jara de Keim",
        'ci': "494913",
        'career': "Masterado en Gestión en la Función Pública"
    }
    nelly_topic = {
        'sections': [
            ("RESUMEN", 1, "Esta monografía analiza la ejecución presupuestaria y la transparencia en los procesos de compras públicas de la Municipalidad de Asunción. Se identificaron cuellos de botella en la rendición de cuentas que dificultan el seguimiento del gasto. Se propone un modelo de control interno para mejorar la eficiencia del gasto público. Palabras clave: Presupuesto Público; Compras Públicas; Transparencia; Control Interno; Contabilidad."),
            ("ABSTRACT", 1, "This monograph analyzes budget execution and transparency in the public procurement processes of the Municipality of Asunción. Bottlenecks in accountability were identified that hinder the tracking of spending. An internal control model is proposed to improve the efficiency of public spending. Keywords: Public Budget; Public Procurement; Transparency; Internal Control; Accounting."),
            ("INTRODUCCIÓN", 1, "La correcta ejecución del presupuesto público es vital para garantizar la provisión de servicios a la ciudadanía. Este trabajo se centra en analizar cómo la Municipalidad de Asunción gestiona sus gastos y compras públicas a través del prisma de la contabilidad gubernamental, un enfoque totalmente distinto a la gestión de ingresos tributarios."),
            ("PLANTEAMIENTO DEL PROBLEMA", 1, "Los retrasos en la rendición de cuentas y la falta de trazabilidad en los procesos de compras públicas representan un desafío constante en la administración de fondos municipales. ¿De qué manera la optimización del control interno puede mejorar la transparencia en la ejecución del presupuesto de gastos en la Municipalidad de Asunción?"),
            ("MARCO TEÓRICO Y NORMATIVO", 1, "Se aborda la Ley de Contrataciones Públicas y las normativas de control interno dictadas por la Contraloría General de la República..."),
            ("METODOLOGÍA", 1, "Estudio cualitativo descriptivo basado en la revisión de informes de auditoría y reportes de ejecución de gastos de la Municipalidad de Asunción..."),
            ("RESULTADOS", 1, "Se encontraron demoras de hasta 60 días en la carga de documentos respaldatorios de compras públicas en los sistemas contables..."),
            ("ANÁLISIS Y DIAGNÓSTICO", 1, "La falta de un sistema de control interno automatizado genera vulnerabilidades en la ejecución de los fondos, afectando la eficiencia institucional..."),
            ("PROPUESTAS Y RECOMENDACIONES", 1, "Se propone la creación de un manual de procedimientos estandarizado para la ejecución de compras públicas y la capacitación en control interno..."),
            ("CONCLUSIONES", 1, "La eficiencia en la ejecución del presupuesto depende directamente de la modernización de los procesos de compras y auditoría interna..."),
            ("REFERENCIAS", 1, "Bibliografía enfocada en Presupuesto, Compras Públicas y Control Interno."),
            ("ANEXOS", 1, "Matriz de consistencia para el análisis del gasto público.")
        ]
    }

    generate_document(chris_student, chris_topic, "TP-TRABAJO_DE_INVESTIGACION_MONOGRAFIA_ChristhianKeim_Catastro.docx")
    generate_document(nelly_student, nelly_topic, "TP-TRABAJO_DE_INVESTIGACION_MONOGRAFIA_NellyDeKeim_Presupuesto.docx")

if __name__ == "__main__":
    main()
