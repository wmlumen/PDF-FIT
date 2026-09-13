import docx
from docx.shared import Pt, Inches, RGBColor, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = OxmlElement(tag)
            for key in ["sz", "val", "color", "space", "shadow"]:
                if key in edge_data:
                    element.set(qn('w:{}'.format(key)), str(edge_data[key]))
            tcBorders.append(element)
    tcPr.append(tcBorders)

def create_template(version=7):
    doc = docx.Document()
    
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Inches(0.5)
    sec.bottom_margin = Inches(0.5)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)

    # 1. Logo
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists('Logo Centuria.png'):
        run = p.add_run()
        run.add_picture('Logo Centuria.png', width=Inches(1.5))
    
    # 2. Institution Name
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("INSTITUTO SUPERIOR EN CIENCIAS EMPRESARIALES")
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.bold = True
    
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run('"CENTURIA"')
    run.font.name = 'Arial'
    run.font.size = Pt(13)
    run.bold = True

    # 3. Main Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run("TRABAJO DE\nINVESTIGACION\nMONOGRAFIA")
    run.font.name = 'Arial'
    run.font.size = Pt(28)
    run.bold = True
    p.paragraph_format.space_after = Pt(12)

    # 4. Student Data Table
    table = doc.add_table(rows=6, cols=3)
    table.autofit = False
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(0.2)
    table.columns[2].width = Inches(4.0)
    
    data = [
        ("MATERIA", ":", "{{ materia }}"),
        ("DOCENTE", ":", "{{ docente }}"),
        ("ALUMNO", ":", "{{ alumno }}"),
        ("CEDULA DE IDENTIDAD N", ":", "{{ cedula }}"),
        ("CARRERA", ":", "{{ carrera }}"),
        ("SECCION / GRUPO", ":", "{{ seccion }}")
    ]
    
    for i, (col1, col2, col3) in enumerate(data):
        row = table.rows[i]
        tr = row._tr
        trHeight = OxmlElement('w:trHeight')
        trHeight.set(qn('w:val'), '450')  # Smaller height
        trHeight.set(qn('w:hRule'), 'exact')
        trPr = OxmlElement('w:trPr')
        trPr.append(trHeight)
        tr.append(trPr)

        for j, text in enumerate([col1, col2, col3]):
            cell = row.cells[j]
            cell.text = text
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j != 1 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(10)
                if j == 0 or j == 1:
                    r.bold = True
                
            set_cell_border(cell, bottom={"sz": 8, "val": "single", "color": "000000"})

    # 5. Uso Interno Section
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    run = p.add_run("DATOS PARA USO INTERNO")
    run.font.name = 'Arial'
    run.font.size = Pt(12)
    run.bold = True

    # 3-column table
    table2 = doc.add_table(rows=2, cols=3)
    table2.style = 'Table Grid'
    for row in table2.rows:
        row.height = Inches(0.2)
    table2.rows[1].height = Inches(0.6)
    
    headers = ["ADMINISTRACION", "COORDINACION", "DOCENTE"]
    for i, h in enumerate(headers):
        cell = table2.cell(0, i)
        cell.text = h
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            r.bold = True

    for i in range(3):
        cell = table2.cell(1, i)
        cell.text = "\n\nFIRMA Y SELLO"
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Arial'
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(100, 100, 100)

    # 2-column table
    table3 = doc.add_table(rows=1, cols=2)
    table3.style = 'Table Grid'
    table3.columns[0].width = Inches(4.5)
    table3.columns[1].width = Inches(1.5)
    
    c1 = table3.cell(0, 0)
    c1.text = "COMPROBANTE DE INGRESO N:"
    p = c1.paragraphs[0]
    p.runs[0].bold = True; p.runs[0].font.size = Pt(9); p.runs[0].font.name = 'Arial'
    
    c2 = table3.cell(0, 1)
    c2.text = "FECHA:"
    p = c2.paragraphs[0]
    p.runs[0].bold = True; p.runs[0].font.size = Pt(9); p.runs[0].font.name = 'Arial'

    doc.add_page_break()

def add_page_number(run):
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = "PAGE"
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

    # Index section (1 inch margins from here on out)
    new_sec = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    new_sec.page_width = Mm(210)
    new_sec.page_height = Mm(297)
    new_sec.top_margin = Inches(1)
    new_sec.bottom_margin = Inches(1)
    new_sec.left_margin = Inches(1)
    new_sec.right_margin = Inches(1)
    
    # Add page numbers to footer
    footer = new_sec.footer
    p_footer = footer.add_paragraph() if not footer.paragraphs else footer.paragraphs[0]
    p_footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run_footer = p_footer.add_run()
    run_footer.font.name = 'Times New Roman'
    run_footer.font.size = Pt(12)
    add_page_number(run_footer)
    
    # Indices
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("INDICE GENERAL")
    run.bold = True; run.font.name = 'Times New Roman'; run.font.size = Pt(12)

    p_toc = doc.add_paragraph()
    run_toc = p_toc.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar'); fldChar3.set(qn('w:fldCharType'), 'end')
    run_toc._r.append(fldChar1); run_toc._r.append(instrText); run_toc._r.append(fldChar2); run_toc._r.append(fldChar3)

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("INDICE DE TABLAS")
    run.bold = True; run.font.name = 'Times New Roman'; run.font.size = Pt(12)

    p_toc2 = doc.add_paragraph()
    run_toc2 = p_toc2.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'TOC \\h \\z \\c "Tabla"'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar'); fldChar3.set(qn('w:fldCharType'), 'end')
    run_toc2._r.append(fldChar1); run_toc2._r.append(instrText); run_toc2._r.append(fldChar2); run_toc2._r.append(fldChar3)

    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("INDICE DE FIGURAS")
    run.bold = True; run.font.name = 'Times New Roman'; run.font.size = Pt(12)

    p_toc3 = doc.add_paragraph()
    run_toc3 = p_toc3.add_run()
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'TOC \\h \\z \\c "Figura"'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar'); fldChar3.set(qn('w:fldCharType'), 'end')
    run_toc3._r.append(fldChar1); run_toc3._r.append(instrText); run_toc3._r.append(fldChar2); run_toc3._r.append(fldChar3)

    doc.add_page_break()
    p = doc.add_paragraph('{{ body_content }}')

    doc.save(f"Plantilla_Centuria_APA{version}.docx")

if __name__ == "__main__":
    create_template(6)
    create_template(7)
