from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
import os

prs = Presentation()

clases_data = [
    {'clase': 1, 'titulo': 'UNIDAD I: Introducción a los Sistemas de Información Computarizados', 'puntos': [
        ("1. Dato e Información", ["Transformar hechos aislados en conocimientos estructurados.", "Significado y valor directo para la toma de decisiones."]), 
        ("2. Tipología de los SI", ["Transición de sistemas transaccionales básicos (TPS).", "Herramientas de apoyo a decisiones (DSS, GDSS, EIS) y sistemas estratégicos (SIS)."]), 
        ("3. Ciclo de Vida", ["Fases evolutivas secuenciales de nacimiento (estudio de factibilidad), desarrollo, operación, mantenimiento y muerte de la tecnología."]),
        ("4. Métodos Alternativos de Adquisición", ["Análisis de conveniencia: método tradicional estructurado vs compra de paquetes.", "Cómputo/desarrollo del usuario final y outsourcing estratégico de TI."])
    ]},
    {'clase': 2, 'titulo': 'UNIDAD II y V: Estrategia y Hardware/Software', 'puntos': [("Estrategia de Negocios", ["Ventajas competitivas y sistemas de información.", "Fuerzas de la industria e impulsos estratégicos."]), ("Sistemas Estratégicos", ["Implantación de sistemas estratégicos.", "Reingeniería de procesos."]), ("Unidad V: Hardware y Software", ["La computadora: definición y componentes básicos.", "Concepto de software."])]},
    {'clase': 3, 'titulo': 'UNIDAD III: Administración de Base de Datos', 'puntos': [("Base de Datos", ["Archivos convencionales vs Definición de base de datos.", "Ventajas en el uso de base de datos."]), ("Manejo del Sistema", ["El DBMS y el Administrador de la base de datos (DBA)."]), ("Modelos", ["Tipos de modelos de base de datos distribuidas.", "Data warehouse."])]},
    {'clase': 4, 'titulo': 'UNIDAD IV: Sistemas Integrados de Gestión', 'puntos': [("Sistemas ERP", ["Sistemas integradores de la administración de empresas (ERP)."]), ("Evaluación", ["Determinación de requerimientos.", "Evaluación técnica y evaluación financiera de las propuestas."]), ("Implantación", ["Actualización y costos de TIC.", "Actividades posteriores a la firma del contrato."])]},
    {'clase': 5, 'titulo': 'UNIDAD VI: Infraestructura de Redes', 'puntos': [("Comunicación de Datos", ["Hardware de apoyo a la comunicación.", "Conectividad y redes computacionales."]), ("Internet e Intranet", ["Dominios y servicios en internet.", "Diferencias entre Intranet y Extranet."]), ("Protocolos", ["Protocolos inalámbricos para internet."])]},
    {'clase': 6, 'titulo': 'UNIDAD VII (Parte 1): Tecnologías DSS y GDSS', 'puntos': [("Sistemas DSS", ["Proceso de toma de decisiones.", "Sistemas de apoyo a las decisiones (DSS). Caso de aplicación."]), ("Sistemas GDSS", ["Sistemas de apoyo para decisiones en grupo (GDSS).", "Ventajas y desventajas del uso de GDSS. Diseño de salas."])]},
    {'clase': 7, 'titulo': 'UNIDAD VII (Parte 2): Inteligencia Artificial', 'puntos': [("Inteligencia Artificial", ["Fundamentos de inteligencia artificial."]), ("Sistemas Expertos", ["Beneficios que genera el uso de sistemas expertos y costos.", "El generador de sistemas expertos o Shell."]), ("Selección", ["Selección de aplicaciones para sistemas expertos."])]},
    {'clase': 8, 'titulo': 'UNIDAD VIII: Sistemas de Apoyo a Ejecutivos (EIS)', 'puntos': [("Concepto", ["Características de los Sistemas de apoyo a ejecutivos (EIS)."]), ("Desarrollo", ["Factores de éxito y proceso de desarrollo.", "Implantación."]), ("Impacto", ["Efecto del EIS en el proceso de planeación y control."])]},
    {'clase': 9, 'titulo': 'UNIDAD IX (Parte 1): Paradigmas en Internet', 'puntos': [("Negocios On Line", ["Introducción a los negocios por internet.", "Elección de tecnologías."]), ("Comercio Electrónico", ["Concepto y categorías.", "Ventajas, problemáticas, sistemas de pago y aspectos legales."])]},
    {'clase': 10, 'titulo': 'UNIDAD IX (Parte 2): E-Government', 'puntos': [("Gobierno Electrónico", ["Concepto de E-goberment y su impacto institucional."]), ("Firma Digital", ["Aplicación y aspectos legales de la Firma digital."])]},
    {'clase': 11, 'titulo': 'UNIDAD X: Futuro de las TIC', 'puntos': [("Nuevas Tendencias", ["Nuevas tendencias tecnológicas aplicadas a los negocios."]), ("Datos y Redes", ["Cubos de datos.", "Bases de datos post-relacionales.", "Redes de alta velocidad."])]},
    {'clase': 12, 'titulo': 'Estudios de Casos en Laboratorio', 'puntos': [("Aplicación", ["Estudios de casos en laboratorio de informática.", "Visión general de las herramientas de apoyo a las decisiones."]), ("Evaluación", ["Talleres para elaboración y presentación de trabajos.", "Preparación del Compilador."])]}
]

title_slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(title_slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "TECNOLOGÍA DE LA INFORMACIÓN Y LA COMUNICACIÓN (ADE18)"
subtitle.text = "Instituto Superior Centuria\nPrograma Analítico Oficial"

for data in clases_data:
    section_layout = prs.slide_layouts[2] 
    slide = prs.slides.add_slide(section_layout)
    slide.shapes.title.text = f"Clase {data['clase']}: {data['titulo']}"
    
    content_layout = prs.slide_layouts[1] 
    slide2 = prs.slides.add_slide(content_layout)
    slide2.shapes.title.text = f"Temario de la Sesión"
    
    tf = slide2.shapes.placeholders[1].text_frame
    tf.clear()
    
    for section_title, bullets in data['puntos']:
        p = tf.add_paragraph()
        p.text = section_title
        p.font.bold = True
        p.font.size = Pt(22)
        p.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
        for bullet in bullets:
            p2 = tf.add_paragraph()
            p2.text = bullet
            p2.level = 1
            p2.font.size = Pt(18)

output_path = os.path.join(os.getcwd(), "Presentacion_Oficial_ADE18.pptx")
prs.save(output_path)
