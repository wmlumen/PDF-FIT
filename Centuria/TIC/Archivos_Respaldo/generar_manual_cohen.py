import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_title(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Arial'
        if level == 1:
            run.font.color.rgb = RGBColor(0, 51, 102) # Azul oscuro
        elif level == 2:
            run.font.color.rgb = RGBColor(0, 102, 204) # Azul claro

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(11)

# PORTADA
titulo = doc.add_paragraph()
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_tit = titulo.add_run("INSTITUTO SUPERIOR CENTURIA\nASIGNATURA: T.I.C. (ADE18)\n\n")
run_tit.bold = True
run_tit.font.size = Pt(18)
run_tit.font.color.rgb = RGBColor(0, 51, 102)

run_sub = titulo.add_run("MANUAL DE ESTUDIO COMPLETO\n")
run_sub.bold = True
run_sub.font.size = Pt(16)

run_bib = titulo.add_run("Basado en el texto: 'Tecnologías de Información en los Negocios' (Cohen & Asín)")
run_bib.italic = True
run_bib.font.size = Pt(12)

doc.add_page_break()

# INTRODUCCIÓN
add_title(doc, "Introducción al Manual", level=1)
doc.add_paragraph("Este manual ha sido desarrollado para acompañar al estudiante a lo largo de las 10 unidades de la asignatura. Todo el contenido teórico, definiciones y frameworks están extraídos y alineados con la bibliografía oficial del curso: el libro 'Tecnologías de Información en los Negocios' de Daniel Cohen Kare y Enrique Asín Lares.")
doc.add_paragraph()

unidades = [
    {
        "titulo": "UNIDAD I: Introducción a los Sistemas de Información",
        "contenido": [
            ("Datos vs. Información (Concepto de Cohen):", "Cohen y Asín enfatizan que el 'dato' puede ser un número, una palabra o una imagen sin contexto. La 'información' son datos que han sido moldeados en una forma significativa y útil para los seres humanos. Los Sistemas de Información (SI) son los encargados de realizar este procesamiento fundamental."),
            ("Tipos de Sistemas de Información:", "1. Sistemas Transaccionales (TPS): Su función principal es la recolección de datos y automatización de tareas operativas (ej. facturación, nómina). Son intensivos en entrada y salida de datos.\n2. Sistemas de Apoyo a las Decisiones (DSS/EIS/GDSS): No ahorran mano de obra, sino que proveen información interactiva para resolver problemas complejos.\n3. Sistemas Estratégicos (SIS): Se desarrollan para lograr ventajas competitivas, cambiando la relación con clientes o proveedores."),
            ("Evolución de los SI (Modelo de Nolan):", "El libro cita las etapas de Richard Nolan para la adopción tecnológica en las empresas: 1) Etapa de Inicio (adquisición de la primera computadora), 2) Contagio (proliferación descontrolada de aplicaciones), 3) Control (centralización por parte del departamento de sistemas), 4) Integración, 5) Administración de Datos, y 6) Madurez (la información es un recurso estratégico corporativo).")
        ]
    },
    {
        "titulo": "UNIDAD II y V: Estrategia de Negocios y Arquitectura TI",
        "contenido": [
            ("La Ventaja Competitiva y las TI:", "Según Cohen, los Sistemas Estratégicos cambian las metas, operaciones o productos de una empresa para superar a la competencia. Utilizan el modelo de las 5 Fuerzas de Michael Porter para analizar cómo la TI puede frenar a nuevos competidores o reducir el poder de los proveedores."),
            ("Impulsos Estratégicos (Modelo de Wiseman):", "Cohen y Asín introducen los 'Impulsos Estratégicos' de Charles Wiseman. Las empresas usan la TI para lograr: Diferenciación (productos únicos), Liderazgo en Costos, Crecimiento, Alianzas (Joint Ventures) o Innovación (crear un mercado nuevo)."),
            ("Hardware y Software (Unidad V):", "El hardware abarca los componentes físicos clasificados por generaciones tecnológicas. El software se divide en Software de Sistema (como Windows o Linux, que administran el hardware) y Software de Aplicación (los programas que usa el usuario, como un ERP o Excel).")
        ]
    },
    {
        "titulo": "UNIDAD III: Fundamentos de Administración de Bases de Datos",
        "contenido": [
            ("El Enfoque de Base de Datos:", "Frente al antiguo sistema de 'Archivos Tradicionales' (donde los datos se repetían en cada departamento), Cohen explica que el DBMS (Database Management System) centraliza los datos, elimina la redundancia y asegura la integridad."),
            ("El Administrador de la Base de Datos (DBA):", "Es la figura humana (profesional técnico) responsable de definir las reglas de acceso, otorgar contraseñas, hacer copias de seguridad (backups) y garantizar que la base de datos esté disponible y segura."),
            ("Data Warehouse y Cubos de Datos:", "Las BD Transaccionales son para el día a día. El Data Warehouse es un almacén masivo de datos históricos diseñado exclusivamente para que los directivos hagan 'Minería de Datos' y descubran patrones ocultos de consumo.")
        ]
    },
    {
        "titulo": "UNIDAD IV: Sistemas Integrados de Gestión (ERP)",
        "contenido": [
            ("Definición de ERP (Enterprise Resource Planning):", "Son sistemas que integran TODO el negocio (Finanzas, Ventas, Producción, RRHH) en una sola plataforma y base de datos (ej. SAP). Cohen destaca que el mayor desafío de un ERP no es técnico, sino humano: obliga a la empresa a cambiar su forma de trabajar (Reingeniería)."),
            ("Evaluación de Propuestas:", "Antes de comprar software costoso, las empresas realizan una 'Evaluación Técnica' (¿El sistema tiene las funciones que necesito?) y una 'Evaluación Financiera' (TIR, VAN y Costo Total de Propiedad).")
        ]
    },
    {
        "titulo": "UNIDAD VI: Infraestructura de Redes y Sistemas Empresariales",
        "contenido": [
            ("Topologías y Redes:", "Cohen detalla cómo las computadoras se conectan en redes LAN (Locales) o WAN (Área Amplia). Se analizan los medios físicos (cobre, fibra óptica) e inalámbricos (microondas, satélite, Wi-Fi)."),
            ("Internet, Intranet y Extranet:", "Internet es la red pública mundial (protocolo TCP/IP). Intranet es una red privada dentro de la empresa que usa tecnología web (solo para empleados). Extranet es cuando la empresa abre parte de su Intranet a proveedores o clientes estratégicos con contraseñas seguras.")
        ]
    },
    {
        "titulo": "UNIDAD VII: Tecnologías de Apoyo a Decisiones (DSS y GDSS)",
        "contenido": [
            ("El Modelo de Toma de Decisiones de H. Simon:", "Cohen explica las 4 fases: Inteligencia (detectar el problema), Diseño (crear soluciones), Selección (elegir una) e Implantación."),
            ("Sistemas DSS y GDSS:", "El DSS (Decision Support System) permite hacer análisis 'What-If' (¿Qué pasaría si subo el precio un 5%?). El GDSS (Group DSS) es para decisiones en equipo; su característica clave es el 'Anonimato', permitiendo que los empleados voten ideas sin miedo a represalias del jefe."),
            ("Sistemas Expertos (Inteligencia Artificial):", "Capturan el conocimiento de un experto humano (ej. un médico) en un 'Motor de Inferencia' usando reglas 'Si-Entonces'. Un 'Shell' es un sistema experto pre-programado y vacío, listo para que se le carguen las reglas de cualquier profesión.")
        ]
    },
    {
        "titulo": "UNIDAD VIII: Sistemas de Apoyo a Ejecutivos (EIS)",
        "contenido": [
            ("Interfaces para la Alta Dirección:", "Los CEOs no tienen tiempo para leer manuales. Según Cohen, un EIS debe tener gráficas de alta calidad, usar colores (semáforos) y permitir el 'Drill-Down' (hacer clic en un gráfico de ventas anuales para desglosarlo en meses, luego en días, luego por vendedor)."),
            ("Factores Críticos de Éxito (FCE):", "Un EIS debe diseñarse monitoreando exclusivamente los FCE, es decir, las 3 o 4 variables que definirán si la empresa sobrevive o quiebra.")
        ]
    },
    {
        "titulo": "UNIDAD IX: Comercio Electrónico (E-Commerce) y E-Government",
        "contenido": [
            ("Categorías del E-Commerce:", "B2C (Empresa a Consumidor, ej. Amazon). B2B (Empresa a Empresa, transacciones masivas). C2C (Consumidor a Consumidor, ej. MercadoLibre/eBay)."),
            ("Firma Digital:", "Para que el E-Commerce y el E-Government (Gobierno Digital) funcionen legalmente, Cohen resalta la 'Firma Digital'. Utiliza criptografía asimétrica (claves públicas y privadas) para garantizar la Autenticidad (saber quién firmó) y el No Repudio (el firmante no puede negarlo).")
        ]
    },
    {
        "titulo": "UNIDAD X: El Futuro de las TIC",
        "contenido": [
            ("Bases de Datos Post-Relacionales y Redes 5G:", "Para manejar el 'Big Data' (datos no estructurados como videos de redes sociales y sensores IoT), las bases de datos tradicionales relacionales ya no son suficientes. El futuro pertenece a bases de datos masivas (NoSQL) y a las redes de alta velocidad que eliminarán los tiempos de carga (latencia).")
        ]
    }
]

for unidad in unidades:
    add_title(doc, unidad["titulo"], level=2)
    for subtitulo, texto in unidad["contenido"]:
        p = doc.add_paragraph()
        p.add_run(subtitulo + " ").bold = True
        p.add_run(texto)
    doc.add_paragraph()

add_title(doc, "Metodología de Estudio Recomendada", level=2)
doc.add_paragraph("Este documento compendia los núcleos teóricos del libro de Daniel Cohen y Enrique Asín. Se recomienda al alumno utilizar este manual para resolver el 'Compilador de Ejercicios' exigido en el programa, ya que las definiciones aquí presentes son las oficiales para la evaluación final.")

output_path = os.path.join(os.getcwd(), "Manual_Estudio_Cohen_Asin_ADE18.docx")
doc.save(output_path)
