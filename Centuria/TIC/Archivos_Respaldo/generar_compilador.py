import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("INSTITUTO SUPERIOR CENTURIA\nCOMPILADOR DE EJERCICIOS OFICIAL")
run.bold = True
run.font.size = Pt(16)
run.font.color.rgb = RGBColor(0, 51, 102)

doc.add_paragraph()
datos = doc.add_paragraph()
datos.add_run("Asignatura: ").bold = True
datos.add_run("Tecnología de la Información y la Comunicación (ADE18)\n")
datos.add_run("Curso: ").bold = True
datos.add_run("Segundo       ")
datos.add_run("Semestre: ").bold = True
datos.add_run("Tercero\n")
datos.add_run("Nombre del Estudiante: ").bold = True
datos.add_run("____________________________________________________\n")
datos.add_run("Cédula de Identidad: ").bold = True
datos.add_run("________________________")

doc.add_paragraph("_" * 70)
intro = doc.add_paragraph()
intro.add_run("Instrucciones: ").bold = True
intro.add_run("El presente documento es un requisito obligatorio para acceder al Examen Final de la asignatura. Debe completarse individualmente aplicando los conceptos teóricos y prácticos de las 10 Unidades del programa oficial.")
intro.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
doc.add_paragraph()

ejercicios = [
    {"clase": "Clase 1", "unidad": "UNIDAD I", "consigna": "Elabore un resumen describiendo el 'Ciclo de vida de los sistemas de información' y explique en qué consiste el método de adquisición por 'Outsourcing'."},
    {"clase": "Clase 2", "unidad": "UNIDAD II y V", "consigna": "Desarrolle un análisis explicando cómo los sistemas de información pueden generar 'Ventajas competitivas' alterando las 'Fuerzas de la industria'. Mencione además la definición y componentes básicos de una computadora."},
    {"clase": "Clase 3", "unidad": "UNIDAD III", "consigna": "Realice un cuadro comparativo entre los 'Archivos convencionales' y las 'Bases de datos', enumerando al menos 3 ventajas de usar una base de datos. Defina el concepto de Data warehouse."},
    {"clase": "Clase 4", "unidad": "UNIDAD IV", "consigna": "Describa cuáles son los pasos para la 'Determinación de requerimientos' y explique la diferencia entre la 'Evaluación técnica de propuestas' y la 'Evaluación financiera de las propuestas'."},
    {"clase": "Clase 5", "unidad": "UNIDAD VI", "consigna": "Defina y establezca las diferencias conceptuales y de uso entre 'Internet', 'Intranet' y 'Extranet'."},
    {"clase": "Clase 6", "unidad": "UNIDAD VII (Parte 1)", "consigna": "Enumere las 'Características de los sistemas de apoyo para la toma de decisiones en grupo (GDSS)' y cite sus principales ventajas y desventajas."},
    {"clase": "Clase 7", "unidad": "UNIDAD VII (Parte 2)", "consigna": "Investigue y defina qué es un 'Sistema experto', indicando los 'Beneficios que genera su uso' y los 'Costos que involucra'. Explique qué es un 'Generador de sistemas expertos o Shell'."},
    {"clase": "Clase 8", "unidad": "UNIDAD VIII", "consigna": "Describa el 'Concepto y Características' de los Sistemas de Apoyo a Ejecutivos (EIS) y explique cuál es su 'Efecto en el proceso de planeación y control de la organización'."},
    {"clase": "Clase 9", "unidad": "UNIDAD IX (Parte 1)", "consigna": "Defina el 'Comercio electrónico', enumere sus categorías, y detalle cuáles son sus principales ventajas y problemáticas según el programa de la asignatura."},
    {"clase": "Clase 10", "unidad": "UNIDAD IX (Parte 2)", "consigna": "Investigue y elabore un ensayo breve sobre el concepto de 'E-goberment' y la importancia legal y técnica de la 'Firma digital'."},
    {"clase": "Clase 11", "unidad": "UNIDAD X", "consigna": "Elabore un resumen explicativo sobre las 'Nuevas tendencias tecnológicas aplicadas a los negocios' abarcando específicamente los 'Cubos de datos', las 'Bases de datos post-relacionales' y las 'Redes de alta velocidad'."},
    {"clase": "Clase 12", "unidad": "Evaluación Final", "consigna": "Revise todo su 'Compilador de Ejercicios'. Asegúrese de haber completado las consignas de las Unidades I a la X estrictamente con los contenidos teóricos desarrollados. Prepárese para presentar este trabajo final."}
]

for i, ej in enumerate(ejercicios):
    p_header = doc.add_paragraph()
    p_header.add_run(f"Evaluación {ej['clase']} | {ej['unidad']}").bold = True
    p_header.style.font.size = Pt(12)
    p_header.style.font.color.rgb = RGBColor(0, 102, 204)
    p_cons = doc.add_paragraph()
    p_cons.add_run("Consigna: ").bold = True
    p_cons.add_run(ej['consigna'])
    p_resp = doc.add_paragraph()
    p_resp.add_run("Respuesta / Desarrollo:\n\n").bold = True
    if "ensayo" in ej['consigna'].lower() or "resumen" in ej['consigna'].lower() or "cuadro" in ej['consigna'].lower():
        for _ in range(25): doc.add_paragraph()
        doc.add_page_break()
    else:
        for _ in range(15): doc.add_paragraph()
        if (i + 1) % 2 == 0: doc.add_page_break()

output_path = os.path.join(os.getcwd(), "Compilador_Ejercicios_ADE18.docx")
doc.save(output_path)
